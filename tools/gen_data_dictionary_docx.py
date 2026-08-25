#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EuroPanel — Data dictionary generator (Word).

Usage:
    python tools/gen_data_dictionary_docx.py

Produces EuroPanel_DataDictionary.docx at the repository root.

This script does not read docs/fields.js itself: it asks
tools/gen_data_dictionary.mjs for the model that file builds (via
`node tools/gen_data_dictionary.mjs --json`), and lays that same model out
in python-docx. That model is already checked, at generation time, against
docs/fields.js, supabase/migrations/002_relational_schema.sql and
supabase/migrations/003_seed_ref_lists.sql (see gen_data_dictionary.mjs) —
duplicating that logic here in Python would be a second place for the two
to quietly disagree.

No markdown-to-Word conversion is used: the Word document is built directly
from the JSON model with python-docx, table by table.

Requires: python-docx (already a project dependency; version 1.2.0 verified
against this script). No other dependency is added.
"""
import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / 'EuroPanel_DataDictionary.docx'
GENERATOR = ROOT / 'tools' / 'gen_data_dictionary.mjs'

# Palette lifted from the previous EuroPanel_DataDictionary.docx, so the new,
# accurate document reads as the same deliverable rather than a different
# artefact the client has to get used to.
NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
HEADER_FILL = '1E3A5F'
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x55, 0x55)
BODY_FONT = 'Calibri'
MONO_FONT = 'Courier New'


def load_model():
    # Captured as bytes and decoded explicitly as UTF-8: on Windows,
    # subprocess.run(..., text=True) decodes with the process's default
    # locale encoding (commonly cp1252), which silently mangles the
    # em-dashes, subscripts (SO₂, NOₓ) and other non-ASCII characters that
    # Node writes as UTF-8. json.loads() would not catch this — cp1252 can
    # decode almost any byte sequence, so the corruption is silent.
    proc = subprocess.run(
        ['node', str(GENERATOR), '--json'],
        cwd=ROOT, capture_output=True, check=False)
    if proc.returncode != 0:
        sys.stderr.write(proc.stderr.decode('utf-8', errors='replace'))
        raise SystemExit(
            f'gen_data_dictionary.mjs --json exited with {proc.returncode}')
    return json.loads(proc.stdout.decode('utf-8'))


# ── Formatting helpers ──────────────────────────────────────────────────

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): hex_color,
    })
    tcPr.append(shd)


def add_title_block(doc, model):
    p = doc.add_paragraph()
    r = p.add_run('EuroPanel — Data Dictionary')
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = BODY_FONT

    p2 = doc.add_paragraph()
    r2 = p2.add_run('Relational schema for the WBP BREF questionnaire')
    r2.font.size = Pt(14)
    r2.font.color.rgb = BLUE
    r2.font.name = BODY_FONT

    doc.add_paragraph()
    p3 = doc.add_paragraph()
    r3 = p3.add_run(
        f"Generated on {model['generatedAt']}  |  Schema: europanel  |  "
        f"{model['counts']['tables']} questionnaire tables + "
        f"{model['counts']['coreTables']} core tables  |  "
        f"{model['counts']['columns']} columns  |  "
        f"{model['counts']['refLists']} reference lists "
        f"({model['counts']['refEntries']} entries)")
    r3.font.size = Pt(10)
    r3.font.color.rgb = GREY
    r3.font.name = BODY_FONT
    doc.add_paragraph()


def add_heading1(doc, text, page_break_before=False):
    p = doc.add_paragraph(style='Heading 1')
    if page_break_before:
        run0 = p.add_run()
        run0.add_break(WD_BREAK.PAGE)
    r = p.add_run(text)
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = BODY_FONT
    return p


def add_heading2(doc, text, mono=False):
    p = doc.add_paragraph(style='Heading 2')
    r = p.add_run(text)
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = MONO_FONT if mono else BODY_FONT
    return p


def add_body(doc, text, italic=False, size=10, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = BODY_FONT
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p


def add_code_block(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = MONO_FONT
    r.font.size = Pt(9)
    return p


def add_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = WHITE
        r.font.name = BODY_FONT
        set_cell_shading(hdr[i], HEADER_FILL)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run('' if value is None else str(value))
            r.font.size = Pt(9)
            r.font.name = BODY_FONT

    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table


def code_join(items):
    return ', '.join(f'`{i}`' for i in items)


# ── Sections ─────────────────────────────────────────────────────────────

def build_purpose(doc, model):
    add_heading1(doc, '1. Purpose')
    add_body(doc,
        'This document is the data dictionary for the EuroPanel WBP BREF '
        'questionnaire database: every field, its meaning, unit (where '
        'applicable), allowed values, and the questionnaire item it '
        'corresponds to. It exists so that a competent third party who has '
        'never seen this system can understand the database — and '
        'extract a complete, structured copy of the data — without '
        'reading the application code.')
    add_body(doc,
        'It is generated, not hand-written, from docs/fields.js — the '
        'same manifest that generates the database schema '
        '(supabase/migrations/002_relational_schema.sql) and the '
        'reference-list seed (supabase/migrations/003_seed_ref_lists.sql). '
        'The document and the schema come from one source and are checked '
        'against each other at generation time; they cannot silently drift '
        'apart. Regenerate after any change to the manifest with:')
    add_code_block(doc, 'node tools/gen_data_dictionary.mjs > EuroPanel_DataDictionary.md\n'
                         'python tools/gen_data_dictionary_docx.py')
    doc.add_paragraph()
    c = model['counts']
    add_body(doc,
        f"Covers {c['tables']} questionnaire tables ({c['columns']} columns), "
        f"{c['coreTables']} core tables, and {c['refLists']} reference lists "
        f"({c['refEntries']} coded entries).")


def build_how_to_read(doc):
    add_heading1(doc, '2. How to read this document')
    add_body(doc,
        'The questionnaire has 12 pages, numbered 0 to 11 (a 13th page, '
        '"Review & Submit", is a read-only summary and holds no data of its '
        'own). Each page maps to one or more tables in the europanel schema:')
    for bullet in [
        'Most sections are a single table with one row per submission — '
        'one column per question asked on that page.',
        'A repeating section — dryers, combustion units, emission '
        'points, and so on — becomes a child table with one row per '
        'instance the operator added. Instances are numbered by idx, '
        'starting at 1, in the order the operator entered them.',
        'A section built around a fixed list of items — pollutants, '
        'fuels, raw materials — becomes a child table with one row per '
        'applicable code, not per submission. Codes come from a reference '
        'list; § 5 lists every code with its label and unit.',
        'Some tables nest inside another questionnaire table rather than '
        'attaching directly to a submission — for example '
        'combustion_unit_fuels, one row per fuel used by a specific '
        'combustion unit. These carry a foreign key to their parent '
        'table’s row instead of (or in addition to) one to the '
        'submission.',
    ]:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(bullet)
        r.font.size = Pt(10)
        r.font.name = BODY_FONT
    doc.add_paragraph()
    add_body(doc,
        'A handful of columns appear on nearly every table below but are '
        'never asked of the operator — they come from the structure of '
        'the form, or from the database itself:')
    add_table(doc, ['Column', 'Meaning'], [
        ('id', 'Surrogate primary key, generated by the database. Every '
                'table has one, except the 1:1 tables described above, '
                'which use submission_id as their primary key instead.'),
        ('submission_id', 'Foreign key to submissions. Present on every '
                'table that is not nested inside another questionnaire '
                'table; identifies which submission the row belongs to.'),
        ('idx', 'Position of a repeating instance within its section '
                '(1, 2, 3, …). Present only on repeating-section '
                'tables.'),
        ('code, list_code', 'code is the fixed item this row describes, '
                'e.g. nox or roundwood; list_code names the reference list '
                'it belongs to, and is the same value for every row of '
                'that table. Present only on fixed-code-group tables.'),
        ('<parent>_id', 'On a table nested inside another questionnaire '
                'table, the foreign key to the specific parent row — '
                'for example combustion_unit_id on combustion_unit_fuels '
                '— used instead of a direct link to the submission.'),
        ('updated_at', 'Timestamp of the last write to the row, maintained '
                'automatically by a database trigger, never set by the '
                'application.'),
    ], col_widths_cm=[4, 12])


def build_core_tables(doc, model):
    add_heading1(doc, '3. Core tables', page_break_before=True)
    add_body(doc,
        f"These {len(model['core'])} tables carry no questionnaire field, "
        'so the manifest has nothing to say about them and they are '
        'described here directly. Two are inherited from the original '
        'schema (migration 001); the other five were introduced with the '
        'relational schema and are written by hand in tools/gen_schema.mjs '
        'rather than generated, because they describe infrastructure — '
        'companies, plants, reporting cycles, reference data, raw '
        'payloads, the audit trail — that the field manifest has no '
        'reason to know about.')
    add_table(doc, ['Table', 'Source', 'Purpose'],
              [(c['table'], c['source'], c['purpose']) for c in model['core']],
              col_widths_cm=[3, 3.5, 10])


def build_pages(doc, model):
    add_heading1(doc, '4. Questionnaire pages', page_break_before=True)
    add_body(doc,
        'One subsection per page, in page order. For each table: what it '
        'is for, its kind (1:1 / repeating / fixed-code group), its parent '
        'if it is nested, and its columns.')
    for page in model['pages']:
        add_heading1(doc, f"Page {page['id']} — {page['title']}",
                     page_break_before=True)
        for table in page['tables']:
            add_heading2(doc, table['name'], mono=True)
            add_body(doc, f"Kind: {table['kindLine']}", italic=True)
            add_body(doc, f"Parent: {table['parentLine']}", italic=True)
            add_body(doc, table['purpose'])
            rows = [(c['column'], c['sqlType'], c['formField'], c['notes'])
                    for c in table['columns']]
            add_table(doc, ['Column', 'SQL type', 'Form field', 'Notes'], rows,
                      col_widths_cm=[3.5, 2.5, 4.5, 6])


def build_reference_lists(doc, model):
    add_heading1(doc, '5. Reference lists', page_break_before=True)
    add_body(doc,
        'Each fixed-code group in § 4 draws its codes from one of the '
        'following lists. This is where a code such as nox or ext_prod_res '
        'gets a human-readable label and, where applicable, a unit.')
    for rl in model['refLists']:
        add_heading2(doc, rl['name'], mono=True)
        add_body(doc, f"Used by: {', '.join(rl['usedBy'])}", italic=True)
        rows = [(e['code'], e['label'], e['unit'] or '—') for e in rl['entries']]
        add_table(doc, ['Code', 'Label', 'Unit'], rows, col_widths_cm=[4, 8, 4])


def build_inconsistencies(doc, model):
    add_heading1(doc, '6. Known inconsistencies in the questionnaire',
                 page_break_before=True)
    for item in model['inconsistencies']:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(item)
        r.font.size = Pt(10)
        r.font.name = BODY_FONT
    doc.add_paragraph()
    add_body(doc, model['inconsistenciesClosing'])


def strip_backticks(text):
    """python-docx renders plain runs; the Markdown backticks used for
    `code font` in the JSON model's prose are stripped here rather than
    rendered literally, since this script does not implement inline
    Markdown formatting."""
    return text.replace('`', '')


def clean_model_strings(obj):
    if isinstance(obj, str):
        return strip_backticks(obj)
    if isinstance(obj, list):
        return [clean_model_strings(x) for x in obj]
    if isinstance(obj, dict):
        return {k: clean_model_strings(v) for k, v in obj.items()}
    return obj


def main():
    model = clean_model_strings(load_model())

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    add_title_block(doc, model)
    build_purpose(doc, model)
    build_how_to_read(doc)
    build_core_tables(doc, model)
    build_pages(doc, model)
    build_reference_lists(doc, model)
    build_inconsistencies(doc, model)

    doc.core_properties.title = 'EuroPanel — Data Dictionary'
    doc.core_properties.subject = 'Relational schema for the WBP BREF questionnaire'

    doc.save(str(OUTPUT))
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
